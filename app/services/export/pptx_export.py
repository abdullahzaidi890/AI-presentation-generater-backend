"""Editable PPTX: real text boxes / title placeholders, bullets, theme colors, speaker notes."""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

from .themes import get_theme, hex_to_rgb

W, H = Inches(13.333), Inches(7.5)


def _rgb(h):
    return RGBColor(*hex_to_rgb(h))


def _bg(slide, theme):
    fill = slide.background.fill
    if len(theme["bg"]) == 1:
        fill.solid()
        fill.fore_color.rgb = _rgb(theme["bg"][0])
    else:
        fill.gradient()
        fill.gradient_angle = 315
        fill.gradient_stops[0].color.rgb = _rgb(theme["bg"][0])
        fill.gradient_stops[1].color.rgb = _rgb(theme["bg"][1])


def _style_run(run, size, color, font, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = _rgb(color)


def _title(slide, text, theme, size=34, top=Inches(0.5), center=False, height=Inches(1.2)):
    ph = slide.shapes.title  # real title placeholder
    ph.left, ph.top, ph.width, ph.height = Inches(0.8), top, W - Inches(1.6), height
    tf = ph.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    _style_run(r, size, theme["title"], theme["font_heading"], bold=True)


def _bullets_box(slide, items, theme, left, top, width, height, size=22, bullets=True):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = item
        _style_run(r, size, theme["text"], theme["font_body"])
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.35))))
            pPr.set("indent", str(-Emu(Inches(0.3))))
            clr = etree.SubElement(pPr, qn("a:buClr"))
            srgb = etree.SubElement(clr, qn("a:srgbClr"))
            srgb.set("val", theme["accent"].lstrip("#"))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")
    return tb


def _accent_bar(slide, theme):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(1.2), Inches(0.07))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(theme["accent"])
    bar.line.fill.background()


def _footer(slide, theme, author):
    """Small 'Created by' line at the bottom-left of every slide."""
    if not author:
        return
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(6), Inches(0.35))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = f"Created by {author}"
    _style_run(r, 10, theme["muted"], theme["font_body"])


def _split(items):
    mid = (len(items) + 1) // 2
    return items[:mid], items[mid:]


def build_pptx(deck: dict, theme_id: str, author: str = "") -> bytes:
    theme = get_theme(theme_id)
    prs = Presentation()
    if author:
        prs.core_properties.author = author
    prs.core_properties.title = deck.get("title", "")[:200]
    prs.slide_width, prs.slide_height = W, H
    title_only = prs.slide_layouts[5]
    for s in deck["slides"]:
        slide = prs.slides.add_slide(title_only)
        _bg(slide, theme)
        layout = s.get("layout", "bullets")
        bullets = s.get("bullets", [])
        if layout in ("title", "closing"):
            _title(slide, s["title"], theme, size=44, top=Inches(2.2), center=True, height=Inches(1.6))
            if bullets:
                _bullets_box(slide, bullets, theme, Inches(2), Inches(4.0), W - Inches(4), Inches(2.5),
                             size=22 if layout == "closing" else 24, bullets=(layout == "closing" and len(bullets) > 1))
                for p in slide.shapes[-1].text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER if not (layout == "closing" and len(bullets) > 1) else PP_ALIGN.LEFT
        elif layout == "quote":
            _title(slide, s["title"], theme, size=26, top=Inches(0.5), height=Inches(0.9))
            quote = bullets[0] if bullets else ""
            tb = _bullets_box(slide, [f"“{quote}”"] if quote else [], theme, Inches(1.5), Inches(2.3), W - Inches(3),
                              Inches(2.5), size=32, bullets=False)
            for p in tb.text_frame.paragraphs:
                for r in p.runs:
                    r.font.italic = True
                    r.font.color.rgb = _rgb(theme["accent"])
            if len(bullets) > 1:
                _bullets_box(slide, ["— " + bullets[1]], theme, Inches(1.5), Inches(5.0), W - Inches(3), Inches(0.8),
                             size=20, bullets=False)
        elif layout == "two_column":
            _title(slide, s["title"], theme)
            _accent_bar(slide, theme)
            a, b = _split(bullets)
            colw = (W - Inches(2.1)) / 2
            _bullets_box(slide, a, theme, Inches(0.8), Inches(2.0), int(colw), Inches(4.8), size=20)
            _bullets_box(slide, b, theme, Inches(0.8) + int(colw) + Inches(0.5), Inches(2.0), int(colw), Inches(4.8), size=20)
        elif layout == "chart":
            _title(slide, s["title"], theme)
            _accent_bar(slide, theme)
            _bullets_box(slide, bullets, theme, Inches(0.8), Inches(2.0), Inches(6.3), Inches(4.8), size=20)
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(2.1), Inches(5.0), Inches(4.2))
            box.fill.background()
            box.line.color.rgb = _rgb(theme["accent"])
            box.line.width = Pt(2)
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = "Suggested visual: " + (s.get("visual_suggestion") or "chart or diagram")
            _style_run(r, 16, theme["muted"], theme["font_body"], italic=True)
        else:
            _title(slide, s["title"], theme)
            _accent_bar(slide, theme)
            _bullets_box(slide, bullets, theme, Inches(0.8), Inches(2.0), W - Inches(1.6), Inches(4.8), size=24)
        _footer(slide, theme, author)
        notes = s.get("speaker_notes", "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
