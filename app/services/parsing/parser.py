"""Safe document parsing. Files are read from memory only (never written to disk or executed)."""
import io
import zipfile
from dataclasses import dataclass

ALLOWED = {".pdf", ".docx", ".txt"}


class DocumentError(Exception):
    """Message is safe to show to the user."""


@dataclass
class ParsedDocument:
    text: str
    truncated: bool
    original_chars: int
    kind: str


def _ext(filename: str) -> str:
    i = filename.rfind(".")
    return filename[i:].lower() if i != -1 else ""


def _truncate(text: str, limit: int) -> tuple[str, bool]:
    """Keep the start (intro) plus evenly spaced chunks so the whole document is represented."""
    if len(text) <= limit:
        return text, False
    head = int(limit * 0.4)
    rest = limit - head
    n_chunks = 6
    size = rest // n_chunks
    span = len(text) - head
    parts = [text[:head]]
    for i in range(n_chunks):
        start = head + int(span * i / n_chunks)
        parts.append("\n[...]\n" + text[start:start + size])
    return "".join(parts), True


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("This PDF is encrypted/password protected. Remove the password and try again.")
        pages = [(p.extract_text() or "") for p in reader.pages]
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError, AssertionError, TypeError):
        raise DocumentError("This PDF looks corrupted and could not be read.") from None
    return "\n".join(pages)


def _read_docx(data: bytes) -> str:
    import docx
    try:
        d = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise DocumentError("This DOCX looks corrupted and could not be read.") from None
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_txt(data: bytes) -> str:
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise DocumentError("This file does not look like a plain text file.")
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise DocumentError("Could not decode the text file. Save it as UTF-8 and try again.")


def parse_document(filename: str, data: bytes, max_bytes: int, max_chars: int) -> ParsedDocument:
    ext = _ext(filename or "")
    if ext not in ALLOWED:
        raise DocumentError("Unsupported file type. Please upload PDF, DOCX or TXT.")
    if len(data) == 0:
        raise DocumentError("The file is empty.")
    if len(data) > max_bytes:
        raise DocumentError(f"The file is too large (limit {max_bytes // (1024 * 1024)} MB).")
    # Check the real content signature, not only the extension.
    if ext == ".pdf":
        if not data.lstrip()[:5] == b"%PDF-":
            raise DocumentError("The file content is not a real PDF.")
        text = _read_pdf(data)
    elif ext == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise DocumentError("The file content is not a real DOCX.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise DocumentError("The file content is not a real DOCX.")
        except zipfile.BadZipFile:
            raise DocumentError("This DOCX looks corrupted and could not be read.") from None
        text = _read_docx(data)
    else:
        text = _read_txt(data)
    text = "\n".join(line.rstrip() for line in text.splitlines())
    if len(text.strip()) < 20:
        hint = " It may be a scanned PDF (images only). Use a text-based file." if ext == ".pdf" else ""
        raise DocumentError("No readable text was found in the file." + hint)
    original = len(text)
    text, truncated = _truncate(text, max_chars)
    return ParsedDocument(text, truncated, original, ext[1:])
